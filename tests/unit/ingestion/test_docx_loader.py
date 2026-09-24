"""DOCX loader tests."""

from pathlib import Path

import pytest
from docx import Document as DOCXDocument

from app.ingestion.errors import DocumentExtractionError
from app.ingestion.loaders.docx_loader import DOCXLoader


def test_docx_loader_extracts_ordered_non_empty_paragraphs(tmp_path: Path) -> None:
    source = tmp_path / "document.docx"
    document = DOCXDocument()
    document.add_paragraph("First paragraph")
    document.add_paragraph("   ")
    document.add_paragraph("Second paragraph")
    document.save(source)

    sections = DOCXLoader().load(source)

    assert [section.text for section in sections] == ["First paragraph", "Second paragraph"]


def test_docx_loader_rejects_invalid_docx(tmp_path: Path) -> None:
    source = tmp_path / "invalid.docx"
    source.write_bytes(b"not a DOCX")

    with pytest.raises(DocumentExtractionError):
        DOCXLoader().load(source)